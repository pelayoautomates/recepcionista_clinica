import io
import logging
import re
from urllib.parse import urljoin, urlparse
from uuid import UUID

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openai import AsyncOpenAI

from config import settings
from database.client import get_supabase
from security import require_admin_key

logger = logging.getLogger(__name__)
router = APIRouter(dependencies=[Depends(require_admin_key)])

MAX_CHARS = 40_000  # límite para no enviar demasiado texto al LLM
MAX_PAGES = 20      # máximo de páginas a rastrear por sitio
CHARS_PER_PAGE = 6_000  # límite de texto por página individual


# ─── Endpoint principal ───────────────────────────────────────────────────────

@router.post("/clinicas/{clinic_id}/configuracion/extraer")
async def extraer_y_generar_prompt(
    clinic_id: UUID,
    url: str | None = Form(default=None),
    archivos: list[UploadFile] = File(default=[]),
):
    """
    Recibe URL y/o archivos de la clínica, extrae el texto,
    llama a GPT-4o para estructurar la info y generar el system prompt.
    """
    textos: list[str] = []

    if url and url.strip():
        try:
            texto_web = await _crawl_site(url.strip())
            textos.append(texto_web)
        except Exception as e:
            logger.warning("No se pudo rastrear %s: %s", url, e)

    for archivo in archivos:
        try:
            texto = await _extract_text_from_file(archivo)
            if texto.strip():
                textos.append(f"[Archivo: {archivo.filename}]\n{texto}")
        except Exception as e:
            logger.warning("No se pudo extraer texto de %s: %s", archivo.filename, e)

    if not textos:
        raise HTTPException(status_code=400, detail="No se pudo extraer texto. Comprueba la URL o el formato del archivo.")

    texto_completo = "\n\n---\n\n".join(textos)[:MAX_CHARS]
    resultado = await _generar_con_ia(str(clinic_id), texto_completo)
    return resultado


@router.post("/clinicas/{clinic_id}/configuracion/guardar")
async def guardar_configuracion(clinic_id: UUID, data: dict):
    """Guarda horarios, servicios y prompt personalizado en la clínica."""
    db = get_supabase()
    campos = {}
    if "prompt_personalizado" in data:
        campos["prompt_personalizado"] = data["prompt_personalizado"]
    if "horarios" in data:
        campos["horarios"] = data["horarios"]
    if "servicios" in data:
        campos["servicios"] = data["servicios"]
    if "notif_webhook" in data:
        campos["notif_webhook"] = data["notif_webhook"]
    if "routing_mode" in data:
        allowed = {"siempre", "fuera_horario", "si_no_contestan"}
        if data["routing_mode"] not in allowed:
            raise HTTPException(status_code=400, detail=f"routing_mode debe ser uno de: {allowed}")
        campos["routing_mode"] = data["routing_mode"]
    if "notification_email" in data:
        campos["notification_email"] = data["notification_email"] or None
    if not campos:
        raise HTTPException(status_code=400, detail="No hay datos para guardar")
    result = db.table("clinicas").update(campos).eq("id", str(clinic_id)).execute()
    return result.data[0]


@router.post("/clinicas/{clinic_id}/onboarding-ok")
async def marcar_onboarding_ok(clinic_id: UUID):
    """Marca el onboarding guiado como completado."""
    db = get_supabase()
    db.table("clinicas").update({"onboarding_ok": True}).eq("id", str(clinic_id)).execute()
    return {"ok": True}


# ─── Extracción de texto ──────────────────────────────────────────────────────

def _html_to_text(html: str) -> str:
    """Convierte HTML a texto plano limpio."""
    # Eliminar bloques no útiles
    html = re.sub(
        r"<(script|style|noscript|nav|footer|header|aside|iframe)[^>]*>.*?</(script|style|noscript|nav|footer|header|aside|iframe)>",
        "", html, flags=re.DOTALL | re.IGNORECASE,
    )
    # Convertir etiquetas de bloque a saltos de línea
    html = re.sub(r"<br\s*/?>", "\n", html, flags=re.IGNORECASE)
    html = re.sub(r"</(p|div|li|h[1-6]|tr|section|article)>", "\n", html, flags=re.IGNORECASE)
    # Quitar resto de etiquetas
    html = re.sub(r"<[^>]+>", " ", html)
    # Entidades HTML comunes
    html = html.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
    html = re.sub(r"&#?\w+;", " ", html)
    # Limpiar espacios
    html = re.sub(r"[ \t]{2,}", " ", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html.strip()[:CHARS_PER_PAGE]


def _extract_links(html: str, current_url: str, base_domain: str) -> list[str]:
    """Extrae enlaces internos del HTML."""
    hrefs = re.findall(r'href=["\']([^"\'#?][^"\']*)["\']', html)
    links: list[str] = []
    skip = ('.pdf', '.jpg', '.jpeg', '.png', '.gif', '.svg', '.css', '.js',
            '.xml', '.ico', '.woff', '.woff2', '.ttf', '.zip', 'mailto:', 'tel:', 'javascript:')
    for href in hrefs:
        full = urljoin(current_url, href)
        parsed = urlparse(full)
        if parsed.netloc == base_domain and not any(full.lower().endswith(s) or s in full for s in skip):
            # Normalise: strip fragment and query
            clean = parsed._replace(fragment="", query="").geturl()
            links.append(clean)
    return links


async def _crawl_site(start_url: str) -> str:
    """Rastrea el sitio web completo (hasta MAX_PAGES páginas) y devuelve el texto concatenado."""
    parsed_start = urlparse(start_url)
    base_domain = parsed_start.netloc

    visited: set[str] = set()
    queue: list[str] = [start_url]
    sections: list[str] = []
    total_chars = 0

    headers = {"User-Agent": "Mozilla/5.0 (compatible; Atiende360Bot/1.0; +https://atiende360.com)"}

    async with httpx.AsyncClient(timeout=20, follow_redirects=True) as client:
        while queue and len(visited) < MAX_PAGES and total_chars < MAX_CHARS:
            url = queue.pop(0)
            if url in visited:
                continue
            visited.add(url)

            try:
                r = await client.get(url, headers=headers)
                r.raise_for_status()
                content_type = r.headers.get("content-type", "")
                if "text/html" not in content_type:
                    continue
            except Exception as exc:
                logger.debug("Omitiendo %s: %s", url, exc)
                continue

            html = r.text
            text = _html_to_text(html)

            if text.strip():
                sections.append(f"[Página: {url}]\n{text}")
                total_chars += len(text)

            # Encolar nuevos enlaces internos
            if len(visited) < MAX_PAGES:
                for link in _extract_links(html, url, base_domain):
                    if link not in visited and link not in queue:
                        queue.append(link)

    logger.info("Crawl completado: %d páginas, %d chars totales", len(visited), total_chars)
    return "\n\n---\n\n".join(sections)[:MAX_CHARS]


async def _extract_text_from_file(archivo: UploadFile) -> str:
    content = await archivo.read()
    name = (archivo.filename or "").lower()

    if name.endswith(".txt") or name.endswith(".csv") or name.endswith(".md"):
        return content.decode("utf-8", errors="ignore")

    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            raise HTTPException(500, "pypdf no instalado — contacta al administrador")

    if name.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            raise HTTPException(500, "python-docx no instalado — contacta al administrador")

    if name.endswith(".xlsx") or name.endswith(".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    fila = [str(c) for c in row if c is not None]
                    if fila:
                        lines.append("\t".join(fila))
            return "\n".join(lines)
        except ImportError:
            # Fallback: try as CSV
            return content.decode("utf-8", errors="ignore")

    # Intento genérico: leer como texto
    return content.decode("utf-8", errors="ignore")


# ─── Generación con IA ────────────────────────────────────────────────────────

_EXTRACTION_PROMPT = """
Eres un asistente experto en configurar recepcionistas virtuales para clínicas.

A partir del siguiente texto sobre una clínica, extrae toda la información relevante y genera:
1. Un resumen estructurado de la clínica
2. Un system prompt completo y profesional para un agente IA que actuará como recepcionista

El system prompt debe:
- Estar escrito en primera persona ("Soy la recepcionista virtual de...")
- Incluir toda la información específica de la clínica (servicios, precios, horarios, ubicación)
- Definir el tono y personalidad del agente
- Incluir instrucciones para gestionar citas, preguntas frecuentes, y cuándo escalar a humano
- Estar en español

Responde EXACTAMENTE en este formato JSON:
{
  "nombre_detectado": "...",
  "resumen": "...",
  "servicios": [{"nombre": "...", "descripcion": "...", "precio": "..."}],
  "horarios": {"lunes": "9:00-18:00", "martes": "...", ...},
  "ubicacion": "...",
  "telefono": "...",
  "web": "...",
  "especialidades": ["...", "..."],
  "faqs": [{"pregunta": "...", "respuesta": "..."}],
  "tono": "profesional/cercano/formal",
  "prompt_generado": "..."
}

Si un campo no se menciona, usa null o array vacío.

TEXTO DE LA CLÍNICA:
"""


async def _generar_con_ia(clinic_id: str, texto: str) -> dict:
    client = AsyncOpenAI(api_key=settings.openai_api_key)

    # Obtener nombre actual de la clínica para contextualizar
    db = get_supabase()
    clinica_res = db.table("clinicas").select("nombre, servicios, horarios").eq("id", clinic_id).single().execute()
    nombre_clinica = clinica_res.data.get("nombre", "la clínica") if clinica_res.data else "la clínica"

    response = await client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {
                "role": "system",
                "content": f"Estás configurando el agente IA para '{nombre_clinica}'. Extrae toda la información útil del texto proporcionado.",
            },
            {
                "role": "user",
                "content": _EXTRACTION_PROMPT + texto,
            },
        ],
        response_format={"type": "json_object"},
        temperature=0.2,
    )

    import json
    try:
        resultado = json.loads(response.choices[0].message.content)
    except Exception:
        raise HTTPException(500, "Error al procesar la respuesta de la IA")

    logger.info("Configuración generada por IA para clínica %s", clinic_id)
    return resultado
